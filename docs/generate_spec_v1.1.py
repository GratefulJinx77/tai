#!/usr/bin/env python3
"""Generate TAI Architecture Spec v1.1 DOCX"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Helper functions
def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
    return h

def add_para(text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = str(val)
    doc.add_paragraph()
    return table

def add_code_block(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.5)
    return p

def add_callout(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(" " + text)
    else:
        p.add_run(text)
    p.paragraph_format.left_indent = Inches(0.3)
    return p

# ============================================================
# COVER PAGE
# ============================================================
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('TAI')
r.font.size = Pt(36)
r.bold = True
r.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run('Team AI Infrastructure')
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)

doc.add_paragraph()

version = doc.add_paragraph()
version.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = version.add_run('Architecture Specification v1.1')
r.font.size = Pt(14)

doc.add_paragraph()

tagline = doc.add_paragraph()
tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tagline.add_run('A discipline framework for AI-assisted development teams')
r.italic = True
r.font.size = Pt(12)

doc.add_paragraph()

origin = doc.add_paragraph()
origin.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = origin.add_run('Forked from PAI (Personal AI Infrastructure) by Daniel Miessler')
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = info.add_run('Jinks Labs\nMarch 2026')
r.font.size = Pt(11)

doc.add_paragraph()

cls = doc.add_paragraph()
cls.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cls.add_run('Classification: Internal')
r.font.size = Pt(10)
r.italic = True

doc.add_page_break()

# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================
add_heading('1. Executive Summary')

doc.add_paragraph(
    'TAI (Team AI Infrastructure) is a forked and adapted version of Daniel Miessler\'s '
    'PAI (Personal AI Infrastructure), restructured for multi-person development teams using '
    'Claude Code as their primary coding tool. Where PAI provides a discipline framework for '
    'individual AI power users, TAI provides an equivalent framework for teams \u2014 ensuring '
    'architectural consistency, verification discipline, and workflow standardization across '
    'all team members.'
)

doc.add_paragraph(
    'TAI is not a replacement for Claude Code. It is the layer on top that makes Claude Code '
    'operate within team-defined boundaries. Claude Code is the engine. TAI is the assembly line.'
)

doc.add_paragraph(
    'TAI is project-agnostic. It ships as a generic framework that any development team can adopt. '
    'Project-specific configuration \u2014 boundary rules, team members, skills, sprint content \u2014 '
    'is defined per-project in .tai/config/. No project-specific references exist in the core framework.'
)

add_callout(
    'AI coding tools generate code faster than teams can verify it. Without a shared discipline '
    'framework, each developer\'s AI produces code that works in isolation but conflicts architecturally. '
    'The result is a codebase that compiles but doesn\'t cohere. TAI enforces coherence at the tool level, '
    'not the policy level.',
    'The Problem TAI Solves:'
)

add_heading('1.1 Design Principles', level=2)

principles = [
    ('Project-scoped, not personal.', 'TAI configuration lives in the repository, not in individual developer environments. Every team member gets the same discipline framework on git pull.'),
    ('Enforced, not suggested.', 'Critical protocols are implemented as hooks and gates that block progress on failure \u2014 not as guidelines that developers may or may not follow.'),
    ('Role-aware.', 'Different team roles need different operational context. An architect\'s session loads differently from a developer\'s session, which loads differently from a QA session.'),
    ('Sprint-aware.', 'The framework knows what sprint the team is in, what the objectives are, and what "done" looks like. This context is injected into every session automatically.'),
    ('Observable.', 'Every session, hook execution, boundary check, and verification gate produces structured telemetry. The CLI status line provides real-time visibility into session health and sprint progress.'),
    ('Composable.', 'Skills are modular, reusable packages. The team skill library grows over time as patterns are codified. Skills from one project can be adapted for another.'),
    ('Versioned.', 'TAI uses semver in a .tai/VERSION manifest. Skill and hook changes are tracked. Breaking changes are communicated clearly.'),
]
for title_text, desc in principles:
    p = doc.add_paragraph()
    r = p.add_run(title_text + ' ')
    r.bold = True
    p.add_run(desc)

# ============================================================
# 2. RELATIONSHIP TO PAI
# ============================================================
doc.add_page_break()
add_heading('2. Relationship to PAI')

doc.add_paragraph(
    'TAI is a fork of PAI (github.com/danielmiessler/Personal_AI_Infrastructure), licensed under MIT. '
    'The fork preserves PAI\'s core architectural concepts while replacing the personal context layer '
    'with a team and project context layer.'
)

add_heading('2.1 What We Keep', level=2)

add_table(
    ['PAI Component', 'TAI Usage', 'Rationale'],
    [
        ['CORE skill (session init)', 'Adapted: loads project + role + sprint context', 'Session initialization is the primary discipline enforcement point'],
        ['Skill system', 'Kept: team-shared skill library', 'Modular, composable workflows are the right abstraction for team patterns'],
        ['Structured response format', 'Kept: standardized AI output', 'Consistent output format accelerates PR review across the team'],
        ['Workflow routing', 'Extended: project-specific routing', 'Auto-directing tasks to correct workflows reduces developer cognitive load'],
        ['Hook system', 'Extended: mandatory verification gates', 'Hooks provide mechanical enforcement, not just policy guidance'],
        ['Observability / JSONL logging', 'Adapted: team-scoped telemetry pipeline', 'Foundation for CLI status line and team metrics'],
        ['GUI installer', 'Adapted: team onboarding wizard', 'PAI\'s guided setup adapted for role-based team configuration'],
        ['Notification system', 'Adapted: Teams/Slack integration', 'PAI uses ntfy/Discord for personal alerts; TAI uses team collaboration platforms'],
    ]
)

add_heading('2.2 What We Strip', level=2)

add_table(
    ['PAI Component', 'Disposition', 'Rationale'],
    [
        ['Personal identity / DA personality', 'Removed', 'Team tool, not personal assistant'],
        ['Life goals and personal preferences', 'Removed', 'Irrelevant to team development context'],
        ['Personal knowledge base retrieval', 'Replaced with project docs', 'Project architecture docs replace personal notes'],
        ['Individual model tier routing', 'Standardized for team', 'Consistent model selection across team members'],
        ['Personal contacts and calendar', 'Removed', 'Out of scope for development infrastructure'],
        ['ElevenLabs TTS / voice system', 'Removed', 'Voice output not applicable to team dev workflow'],
    ]
)

add_heading('2.3 What We Add', level=2)

add_table(
    ['New Component', 'Description', 'Why PAI Doesn\'t Have It'],
    [
        ['Role-based session init', 'Different context loaded per team role', 'PAI has one user; TAI has multiple roles'],
        ['Sprint context injection', 'Current sprint objectives + ISC loaded automatically', 'PAI has no sprint awareness'],
        ['Mandatory verification hooks', 'Test, type-check, boundary gates that block progress', 'PAI trusts the individual; TAI enforces team standards'],
        ['Shared skill library in repo', 'Skills committed to Git, versioned with codebase', 'PAI skills are personal; TAI skills are shared'],
        ['Architecture boundary enforcement', 'Import analysis preventing cross-module violations', 'PAI has no architectural governance concept'],
        ['CLI status line', 'Real-time session health, role, sprint progress in terminal', 'PAI\'s status line is personal; TAI\'s is team-scoped'],
        ['PR preparation automation', 'Auto-generates PR description with AI usage disclosure', 'PAI doesn\'t have a team PR workflow'],
        ['Semver versioning', '.tai/VERSION manifest tracks framework changes', 'PAI evolves informally; TAI needs team coordination'],
    ]
)

# ============================================================
# 3. SYSTEM ARCHITECTURE
# ============================================================
doc.add_page_break()
add_heading('3. System Architecture')

doc.add_paragraph(
    'TAI operates as a configuration, skill, and observability layer within the project repository. '
    'It is not a separate application \u2014 it is a directory structure, set of conventions, and a '
    'CLI status line that Claude Code reads, enforces, and reports through.'
)

add_heading('3.1 Directory Structure', level=2)

add_code_block(
    '.tai/                              # TAI root (committed to repo)\n'
    '\u251c\u2500\u2500 CORE.md                        # Session initialization protocol\n'
    '\u251c\u2500\u2500 VERSION                        # Semver manifest (e.g., 1.0.0)\n'
    '\u251c\u2500\u2500 config/\n'
    '\u2502   \u251c\u2500\u2500 team.yaml                  # Team roles and members\n'
    '\u2502   \u251c\u2500\u2500 project.yaml               # Project metadata and stack\n'
    '\u2502   \u2514\u2500\u2500 models.yaml                # Model selection per task type\n'
    '\u251c\u2500\u2500 context/\n'
    '\u2502   \u251c\u2500\u2500 architecture.md            # System architecture reference\n'
    '\u2502   \u251c\u2500\u2500 boundaries.md              # Hard architectural boundaries\n'
    '\u2502   \u251c\u2500\u2500 patterns.md                # Code patterns and conventions\n'
    '\u2502   \u2514\u2500\u2500 sprint-current.md          # Current sprint objectives + ISC\n'
    '\u251c\u2500\u2500 skills/\n'
    '\u2502   \u251c\u2500\u2500 development/               # Feature development workflows\n'
    '\u2502   \u251c\u2500\u2500 quality/                   # Testing and verification workflows\n'
    '\u2502   \u251c\u2500\u2500 architecture/              # Architecture review workflows\n'
    '\u2502   \u2514\u2500\u2500 project/                   # Project-specific workflows\n'
    '\u251c\u2500\u2500 hooks/\n'
    '\u2502   \u251c\u2500\u2500 pre-commit.sh              # Lint, type-check, boundary verify\n'
    '\u2502   \u251c\u2500\u2500 pre-push.sh                # Full test suite\n'
    '\u2502   \u2514\u2500\u2500 post-session.sh            # Session summary \u2192 JSONL telemetry\n'
    '\u251c\u2500\u2500 roles/\n'
    '\u2502   \u251c\u2500\u2500 architect.md               # Architect session context\n'
    '\u2502   \u251c\u2500\u2500 developer.md               # Developer session context\n'
    '\u2502   \u2514\u2500\u2500 qa.md                      # QA/Test session context\n'
    '\u251c\u2500\u2500 telemetry/\n'
    '\u2502   \u251c\u2500\u2500 sessions.jsonl             # Session activity log (append-only)\n'
    '\u2502   \u251c\u2500\u2500 hooks.jsonl                # Hook execution log\n'
    '\u2502   \u251c\u2500\u2500 boundaries.jsonl           # Boundary check results\n'
    '\u2502   \u2514\u2500\u2500 ai-usage.jsonl             # Model/token usage per session\n'
    '\u2514\u2500\u2500 templates/\n'
    '    \u251c\u2500\u2500 pr-description.md          # PR template with AI disclosure\n'
    '    \u251c\u2500\u2500 sprint-isc.md              # ISC definition template\n'
    '    \u2514\u2500\u2500 boundaries-example.md      # Example boundary config'
)

add_heading('3.2 Session Initialization Flow', level=2)

doc.add_paragraph(
    'When a team member launches Claude Code in the project directory, the following '
    'initialization sequence executes:'
)

steps = [
    ('1. CLAUDE.md loaded', '(standard Claude Code behavior) \u2014 provides base project rules.'),
    ('2. CORE.md loaded', '\u2014 TAI session protocol activates. Reads team.yaml and identifies the current developer\'s role via git user.email match (with TAI_ROLE env var override).'),
    ('3. Role context loaded', '\u2014 architect.md, developer.md, or qa.md injected based on role. Defines operational focus and available skills.'),
    ('4. Sprint context loaded', '\u2014 sprint-current.md injected. Current objectives, ISC, and constraints are immediately available.'),
    ('5. Architecture context loaded', '\u2014 boundaries.md and patterns.md injected. Hard constraints and code conventions are active.'),
    ('6. Hooks registered', '\u2014 pre-commit, pre-push, and post-session hooks activated for the session.'),
    ('7. Telemetry initialized', '\u2014 session start event written to sessions.jsonl. All subsequent hook executions and boundary checks are logged.'),
    ('8. Status line activated', '\u2014 CLI status line displays current role, sprint, ISC progress, and session duration.'),
]
for label, desc in steps:
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(desc)

add_heading('3.3 Role Identification', level=2)

doc.add_paragraph(
    'TAI identifies the current developer\'s role using a two-layer mechanism:'
)

p = doc.add_paragraph()
r = p.add_run('Primary: Git email match. ')
r.bold = True
p.add_run('CORE.md reads git config user.email and matches it against team.yaml entries. Each team member\'s email is mapped to a role. This requires no manual setup \u2014 it works automatically on every session.')

p = doc.add_paragraph()
r = p.add_run('Override: TAI_ROLE environment variable. ')
r.bold = True
p.add_run('A developer can set TAI_ROLE=architect (or developer, qa) in their shell profile to override the git email lookup. This is the escape hatch for scenarios like pair programming, role rotation, or testing.')

doc.add_paragraph(
    'If neither method resolves a role, CORE.md defaults to the "developer" role and logs a warning to telemetry.'
)

# ============================================================
# 4. ROLE SYSTEM
# ============================================================
doc.add_page_break()
add_heading('4. Role System')

doc.add_paragraph(
    'TAI defines three base roles. Each role receives different context at session initialization, '
    'has access to different skills, and operates under different verification requirements. '
    'Projects may define additional roles in team.yaml.'
)

add_heading('4.1 Role Definitions', level=2)

add_table(
    ['Attribute', 'Architect', 'Developer', 'QA / Test'],
    [
        ['Primary focus', 'System coherence, architecture decisions, critical code', 'Feature implementation, pattern adherence', 'Test strategy, coverage, scenario validation'],
        ['Code access', 'Full codebase, all modules', 'Assigned feature scope + shared modules', 'Full codebase (read), test directories (write)'],
        ['PR authority', 'Reviews and approves all PRs', 'Submits PRs for review', 'Submits test PRs, reviews test coverage'],
        ['Skills loaded', 'Architecture, review, planning, all dev skills', 'Development, quality (pre-PR only)', 'Quality, test generation, scenario runner'],
        ['Verification level', 'Self-review + CI', 'Mandatory pre-PR + architect review', 'Mandatory test pass + architect review'],
        ['Architecture decisions', 'Makes and documents decisions', 'Follows established decisions', 'Validates decisions through testing'],
    ]
)

add_heading('4.2 Role Context Files', level=2)

doc.add_paragraph('Each role\'s .md file contains:')

items = [
    'Operational directives \u2014 what this role is responsible for and what it should prioritize.',
    'Boundary constraints \u2014 what this role must NOT do (e.g., developers must not make architecture decisions without consulting the architect).',
    'Skill manifest \u2014 which skills from the team library are available to this role.',
    'Verification checklist \u2014 what must be true before this role\'s work is considered complete.',
    'Escalation protocol \u2014 when and how to escalate decisions to the architect.',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

# ============================================================
# 5. SKILL LIBRARY
# ============================================================
doc.add_page_break()
add_heading('5. Skill Library')

doc.add_paragraph(
    'Skills are modular, composable workflow definitions stored in the .tai/skills/ directory. '
    'Each skill is a markdown file that Claude Code loads when the corresponding workflow is invoked. '
    'Skills define step-by-step procedures, verification checkpoints, and expected outputs.'
)

add_heading('5.1 Core Skills (Ship with TAI)', level=2)

add_table(
    ['Skill', 'Category', 'Description', 'Available To'],
    [
        ['new-endpoint', 'Development', 'Scaffold API route: auth, validation, SQL, integration test, type registration', 'Architect, Developer'],
        ['new-component', 'Development', 'Scaffold UI component: server default, styling, prop types', 'Architect, Developer'],
        ['new-migration', 'Development', 'Numbered SQL migration with sequence check, conflict detection, up/down test', 'Architect, Developer'],
        ['formula-work', 'Development', 'Formula/calculation protocol: edge case checklist, cache implications, test matrix', 'Architect, Developer'],
        ['pre-pr', 'Quality', 'Full verification: lint \u2192 type-check \u2192 boundary \u2192 test \u2192 build \u2192 PR description', 'All roles'],
        ['security-check', 'Quality', 'Role matrix compliance, field-level filtering, unauthorized access probes', 'Architect, QA'],
        ['test-generation', 'Quality', 'Tests from specs/ISC: happy path, failure modes, boundary cases, role scenarios', 'Developer, QA'],
        ['review-pr', 'Architecture', 'PR review protocol: pattern check, boundary audit, dependency review, AI assessment', 'Architect'],
        ['boundary-audit', 'Architecture', 'Cross-module import analysis, detect unauthorized data access paths', 'Architect'],
        ['sprint-planning', 'Architecture', 'ISC definition, task breakdown, dependency mapping, risk identification', 'Architect'],
        ['sprint-init', 'Project', 'Update sprint-current.md, distribute objectives, reset sprint telemetry', 'Architect'],
    ]
)

add_heading('5.2 Skill Anatomy', level=2)

doc.add_paragraph('Every skill follows a standard structure:')

add_code_block(
    '# Skill: New API Endpoint\n\n'
    '## Trigger: Developer requests new API route or endpoint\n\n'
    '## Pre-conditions:\n'
    '- Route does not already exist (check routes directory)\n'
    '- Required types defined or identified for creation\n\n'
    '## Steps:\n'
    '1. Read 2-3 existing route handlers for current patterns\n'
    '2. Confirm role-based access requirements with developer\n'
    '3. Create validation schema\n'
    '4. Create route handler with auth middleware and parameterized SQL\n'
    '5. Register route in app entry point\n'
    '6. Create/update shared types if needed\n'
    '7. Write integration test: happy path, auth rejection, validation failure\n'
    '8. Run test suite (must pass)\n'
    '9. Run type checker (must pass)\n'
    '10. Log skill completion to telemetry\n\n'
    '## Verification Gate:\n'
    '- [ ] Validation schema covers all inputs\n'
    '- [ ] Auth middleware checks correct role(s)\n'
    '- [ ] SQL uses parameterized queries only\n'
    '- [ ] Error response follows standard format\n'
    '- [ ] Integration test covers 3+ scenarios\n'
    '- [ ] Shared types registered\n'
    '- [ ] Test suite passes, type checker passes'
)

# ============================================================
# 6. ENFORCEMENT HOOKS
# ============================================================
doc.add_page_break()
add_heading('6. Enforcement Hooks')

doc.add_paragraph(
    'Hooks are the mechanical enforcement layer. They run automatically at defined trigger points '
    'and block progress if checks fail. Unlike CLAUDE.md rules (which Claude Code usually follows '
    'but can skip under pressure), hooks execute as shell commands that return exit codes. Non-zero '
    'exit = blocked.'
)

doc.add_paragraph(
    'Hook crash policy: if a hook script crashes (not fails \u2014 crashes), it is treated as a failure '
    'and the action is blocked. This conservative approach ensures that broken enforcement infrastructure '
    'is fixed immediately rather than silently bypassed.'
)

add_heading('6.1 Hook Registry', level=2)

add_table(
    ['Hook', 'Trigger', 'Checks', 'On Failure'],
    [
        ['pre-commit', 'Before any git commit', 'Linter, type checker, import boundary scan', 'Commit blocked. Errors displayed.'],
        ['pre-push', 'Before pushing to remote', 'Full test suite', 'Push blocked. Failed tests listed.'],
        ['pre-pr', 'Before PR creation (slash command)', 'All pre-commit + pre-push + build + PR desc validation', 'PR creation blocked until all pass.'],
        ['post-session', 'On Claude Code session end', 'Generate session summary \u2192 JSONL telemetry', 'Advisory only (never blocks).'],
    ]
)

add_heading('6.2 Architecture Boundary Hook', level=2)

doc.add_paragraph(
    'The import boundary scan is TAI\'s most critical hook. It statically analyzes import statements '
    'to enforce service separation. Boundary rules are defined in .tai/context/boundaries.md and are '
    'project-specific.'
)

add_code_block(
    '# .tai/context/boundaries.md (project-specific example)\n\n'
    'FORBIDDEN IMPORTS:\n'
    '  src/services/extraction/** \u2192 src/services/pricing/**\n'
    '  src/services/pricing/** \u2192 src/services/extraction/**\n'
    '  src/services/extraction/** \u2192 src/models/rate_cards.*\n\n'
    'ENFORCEMENT: grep-based static analysis on staged files.\n'
    'If any staged file introduces a forbidden import path, commit is blocked.\n'
    'Exit code 1 with clear message identifying the violation.\n'
    'Violation logged to .tai/telemetry/boundaries.jsonl.'
)

add_heading('6.3 Hook Bypass Protocol', level=2)

doc.add_paragraph(
    'All hooks support bypass via --no-verify. However, TAI intercepts bypass attempts and requires '
    'a mandatory reason string. The bypass and reason are both logged to telemetry.'
)

doc.add_paragraph(
    'This design avoids the "wall of shame" problem: legitimate bypasses (false positives, emergency '
    'hotfixes) are distinguishable from discipline erosion because every bypass has a documented reason. '
    'The architect reviews bypass reasons, not just bypass counts.'
)

add_code_block(
    '# Bypass flow:\n'
    '$ git commit --no-verify -m "fix typo"\n'
    '> TAI: Hook bypass detected. Reason required.\n'
    '> Reason: false positive - eslint rule conflict with generated types\n'
    '> Bypass logged to .tai/telemetry/hooks.jsonl'
)

add_heading('6.4 Telemetry Emission', level=2)

doc.add_paragraph(
    'Every hook execution emits a structured event to the appropriate JSONL log. '
    'Telemetry files are committed to the repository and pushed with code. '
    'A GitHub Action triggers on push to .tai/telemetry/** to aggregate metrics '
    'into summary JSON. Git is the transport \u2014 no separate sync infrastructure required.'
)

add_code_block(
    '{\n'
    '  "timestamp": "2026-03-18T14:23:01Z",\n'
    '  "event": "hook.pre-commit",\n'
    '  "user": "developer-1",\n'
    '  "role": "developer",\n'
    '  "result": "pass | fail | bypass",\n'
    '  "bypass_reason": null,\n'
    '  "details": {\n'
    '    "checks_run": ["eslint", "tsc", "boundary-scan"],\n'
    '    "failures": [],\n'
    '    "duration_ms": 3200,\n'
    '    "files_staged": 4\n'
    '  }\n'
    '}'
)

add_heading('6.5 Telemetry Pipeline', level=2)

doc.add_paragraph(
    'Telemetry flows through git, not a separate infrastructure:'
)

steps_telem = [
    ('1. Collection.', 'Every hook execution, session start/end, boundary check, and skill invocation writes a structured JSON event to .tai/telemetry/*.jsonl on the developer\'s machine.'),
    ('2. Commit.', 'Telemetry files are committed alongside code changes. The post-session hook auto-stages telemetry files.'),
    ('3. Aggregation.', 'A GitHub Action triggers on push to .tai/telemetry/**. It reads raw JSONL, computes aggregate metrics (per-sprint, per-developer, per-day), and commits summary JSON back to the repo.'),
    ('4. Consumption.', 'Claude Code slash commands (/tai-sessions, /tai-bypasses) read telemetry JSONL directly. The CLI status line reads session-level data in real time.'),
]
for label, desc in steps_telem:
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(desc)

# ============================================================
# 7. SPRINT CONTEXT SYSTEM
# ============================================================
doc.add_page_break()
add_heading('7. Sprint Context System')

doc.add_paragraph(
    'TAI maintains awareness of the current sprint through a dedicated context file that is '
    'updated at the start of each sprint and loaded into every session.'
)

add_heading('7.1 Sprint Context File', level=2)

doc.add_paragraph('Template: .tai/context/sprint-current.md')

add_code_block(
    '# Sprint [N]: [Sprint Name]\n\n'
    '## Dates: [start] to [end]\n\n'
    '## Objectives:\n'
    '- [Objective 1]\n'
    '- [Objective 2]\n'
    '- [Objective 3]\n\n'
    '## ISC (Ideal State Criteria):\n'
    '- [ ] [Criterion 1]\n'
    '- [ ] [Criterion 2]\n'
    '- [ ] [Criterion 3]\n'
    '- [ ] All existing tests still pass\n\n'
    '## Task Assignments:\n'
    '- Architect: [scope]\n'
    '- Developer: [scope]\n'
    '- QA: [scope]'
)

# ============================================================
# 8. CLI STATUS LINE
# ============================================================
doc.add_page_break()
add_heading('8. CLI Status Line')

doc.add_paragraph(
    'PAI provides a personal status bar in the terminal \u2014 showing context, weather, versions, '
    'and system health for a single user. TAI adapts this concept for team-scoped visibility: '
    'a persistent CLI status line at the bottom of the Claude Code terminal that displays session '
    'health, role context, and sprint progress at a glance.'
)

add_callout(
    'The CLI status line is not a dashboard. It does not replace project management tools. It answers '
    'one question: Is my current session operating within team discipline? It surfaces the signals that '
    'indicate whether the right role is loaded, the sprint context is current, hooks are active, and the '
    'session is healthy.',
    'Design Philosophy:'
)

add_heading('8.1 Status Line Content', level=2)

doc.add_paragraph(
    'The status line renders a single persistent bar at the bottom of the terminal with the following segments:'
)

add_table(
    ['Segment', 'Data Source', 'Display Example'],
    [
        ['Role', 'team.yaml + git email / TAI_ROLE', 'ARCHITECT'],
        ['Sprint', 'sprint-current.md', 'Sprint 21: AI Ingestion'],
        ['ISC Progress', 'sprint-current.md (parsed)', '5/8 ISC'],
        ['Hook Status', 'Last hook result from hooks.jsonl', 'HOOKS: PASS'],
        ['Session Duration', 'Session start timestamp', '42m'],
        ['TAI Version', '.tai/VERSION', 'v1.0.0'],
        ['Boundary Status', 'boundaries.jsonl (session)', '0 violations'],
    ]
)

add_code_block(
    '# Status line rendering example:\n'
    '\u2502 ARCHITECT \u2502 Sprint 21: AI Ingestion \u2502 5/8 ISC \u2502 HOOKS: PASS \u2502 42m \u2502 TAI v1.0.0 \u2502'
)

add_heading('8.2 Status Line Behavior', level=2)

behavior = [
    ('Refresh.', 'The status line updates after every hook execution, skill invocation, or telemetry write. It reads directly from local JSONL files and sprint-current.md.'),
    ('Color coding.', 'Hook status shows green (PASS), red (FAIL), or yellow (BYPASS). ISC progress shows green (>75%), yellow (25-75%), red (<25%).'),
    ('Error surfacing.', 'If the last hook failed, the status line shows the failure type (e.g., "HOOKS: FAIL tsc") to keep it visible until fixed.'),
    ('Minimal footprint.', 'The status line is a single line. It does not consume vertical space beyond what PAI\'s status bar already uses.'),
]
for label, desc in behavior:
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(desc)

add_heading('8.3 Team Metrics via Slash Commands', level=2)

doc.add_paragraph(
    'Team-wide metrics that would traditionally require a dashboard are surfaced through Claude Code '
    'slash commands implemented as TAI skills. These replace the need for a separate web application.'
)

add_table(
    ['Command', 'Description', 'Access'],
    [
        ['/tai-sessions', 'Session activity: who worked today, session duration, skills used', 'Architect (all), Developer (own)'],
        ['/tai-bypasses', 'Hook bypass log with reasons, filterable by sprint/developer', 'Architect'],
        ['/tai-boundaries', 'Boundary violation history: blocked commits, violation details', 'Architect'],
        ['/tai-sprint', 'Sprint ISC progress, task assignments, days remaining', 'All roles'],
        ['/tai-health', 'Code health summary: test count trend, type safety, dependency audit', 'Architect, QA'],
        ['/tai-usage', 'AI usage metrics: model split, token spend estimate, session counts', 'Architect (all), Developer (own)'],
    ]
)

doc.add_paragraph(
    'Privacy model: developers see their own metrics only. The architect sees aggregate team data. '
    'No developer can see another developer\'s individual metrics. Hook bypass reasons are visible to '
    'the architect because bypasses affect team code quality.'
)

add_heading('8.4 Notification System', level=2)

doc.add_paragraph(
    'TAI posts automated notifications to a team collaboration channel (Teams, Slack, or Discord) '
    'via an incoming webhook configured in team.yaml. Notifications are structured, concise, and actionable.'
)

add_table(
    ['Event', 'Notification', 'Priority'],
    [
        ['PR ready for review', '"@Architect: PR #42 ready for review (3 files, 2 new tests)"', 'Normal'],
        ['Sprint ISC item completed', '"ISC: Document upload \u2014 verified and passing"', 'Normal'],
        ['Boundary violation blocked', '"Boundary violation: developer-1 attempted forbidden import in file.ts"', 'High'],
        ['Hook bypass detected', '"Hook bypass: developer-1 used --no-verify. Reason: [reason]"', 'High'],
        ['All sprint ISC complete', '"Sprint 21: All ISC criteria met. Ready for stakeholder review."', 'High'],
        ['Build failure on main', '"CI failure on main: type-check failed after merge of PR #43"', 'Critical'],
    ]
)

add_heading('8.4.1 Notification Routing', level=3)

add_table(
    ['Recipient', 'Receives'],
    [
        ['Team channel', 'Sprint ISC completions, build failures on main, sprint completion'],
        ['Architect (DM)', 'PR review requests, boundary violations, hook bypasses, dependency changes'],
        ['Developer (DM)', 'PR review feedback, their own hook failures'],
        ['PM (channel)', 'Sprint ISC summary (daily), sprint completion, blocker alerts'],
    ]
)

# ============================================================
# 9. CONFIGURATION
# ============================================================
doc.add_page_break()
add_heading('9. Configuration')

add_heading('9.1 team.yaml', level=2)

add_code_block(
    'team:\n'
    '  name: [Project Name]\n'
    '  members:\n'
    '    - name: [Architect Name]\n'
    '      email: architect@example.com\n'
    '      role: architect\n'
    '      github: [github-username]\n'
    '      skills: [all]\n'
    '    - name: [Developer Name]\n'
    '      email: developer@example.com\n'
    '      role: developer\n'
    '      github: [github-username]\n'
    '      skills: [development, quality]\n'
    '    - name: [QA Name]\n'
    '      email: qa@example.com\n'
    '      role: qa\n'
    '      github: [github-username]\n'
    '      skills: [quality, test-generation]\n\n'
    '  notifications:\n'
    '    platform: teams | slack | discord\n'
    '    webhook_url: https://...\n'
    '    channel: [channel-name]\n'
    '    dm_routing:\n'
    '      pr_review: architect\n'
    '      boundary_violation: architect\n'
    '      build_failure: [architect, developer]'
)

add_heading('9.2 models.yaml', level=2)

add_code_block(
    'models:\n'
    '  default: claude-sonnet-4-6\n'
    '  architecture_decisions: claude-opus-4-6\n'
    '  code_generation: claude-sonnet-4-6\n'
    '  test_generation: claude-sonnet-4-6\n'
    '  pr_review: claude-opus-4-6\n'
    '  documentation: claude-sonnet-4-6'
)

add_heading('9.3 VERSION', level=2)

doc.add_paragraph(
    'TAI uses semver in a .tai/VERSION file. The version is bumped when skills, hooks, or '
    'configuration schema change. The status line displays the current version. CORE.md can '
    'optionally check VERSION on session init and warn if the local checkout is behind.'
)

add_code_block(
    '# .tai/VERSION\n'
    '1.0.0'
)

# ============================================================
# 10. IMPLEMENTATION PLAN
# ============================================================
doc.add_page_break()
add_heading('10. Implementation Plan')

doc.add_paragraph(
    'TAI is implemented in four phases. Phases 1\u20134 ship over one weekend, ready for team launch Monday morning.'
)

add_heading('10.1 Phase Schedule', level=2)

add_table(
    ['Phase', 'Timing', 'Deliverables', 'Hours'],
    [
        ['1: Fork and Strip', 'Saturday AM', 'Clean fork, personal modules removed, directory structure established', '3'],
        ['2: Adapt CORE', 'Saturday PM', 'CORE.md for team context. Role files. Sprint context template. team.yaml configured. Role ID via git email + TAI_ROLE override.', '4'],
        ['3: Build Skills + Hooks', 'Sunday AM', '11 core skills from CLAUDE.md patterns. Verification gates. Hooks implemented and tested. Bypass-with-reason protocol.', '5'],
        ['4: Telemetry + CLI + Integration', 'Sunday PM', 'JSONL logging pipeline. Slash command query skills. CLI status line. GitHub Action for aggregation. Full walkthrough as each role. Fix friction.', '4'],
    ]
)

doc.add_paragraph('Weekend effort (Phases 1\u20134): 16 hours. Ships Monday with full framework.')

add_heading('10.2 Integration with Existing CLAUDE.md', level=2)

doc.add_paragraph(
    'TAI does not replace the project CLAUDE.md \u2014 it extends it. Architecture boundaries, '
    'code patterns, and sprint context are extracted from CLAUDE.md into .tai/context/ files. '
    'CLAUDE.md retains base project rules and adds a single reference line pointing to .tai/.'
)

doc.add_paragraph(
    'The initialization sequence is: CLAUDE.md (project rules) \u2192 .tai/CORE.md (team protocol) '
    '\u2192 role context \u2192 sprint context. Each layer builds on the previous one. A developer '
    'without TAI installed still gets the CLAUDE.md guardrails. TAI adds the discipline enforcement '
    'and observability on top.'
)

add_heading('10.3 CLAUDE.md Migration', level=2)

doc.add_paragraph(
    'Content that moves from CLAUDE.md to .tai/:'
)

add_table(
    ['Content', 'From', 'To'],
    [
        ['Architecture boundaries', 'CLAUDE.md inline rules', '.tai/context/boundaries.md'],
        ['Code patterns / conventions', 'CLAUDE.md inline rules', '.tai/context/patterns.md'],
        ['Sprint objectives / ISC', 'CLAUDE.md or separate doc', '.tai/context/sprint-current.md'],
        ['System architecture reference', 'CLAUDE.md or README', '.tai/context/architecture.md'],
    ]
)

doc.add_paragraph(
    'Content that stays in CLAUDE.md: project name, tech stack summary, build/run commands, '
    'directory overview, and a one-liner: "See .tai/ for team framework, roles, skills, and sprint context."'
)

# ============================================================
# 11. RISKS AND MITIGATIONS
# ============================================================
doc.add_page_break()
add_heading('11. Risks and Mitigations')

add_table(
    ['Risk', 'Impact', 'Likelihood', 'Mitigation'],
    [
        ['Upstream PAI breaking changes', 'Medium', 'Low', 'Fork is pinned. No auto-update. Selective cherry-pick only.'],
        ['Context window consumption', 'High', 'Medium', 'Measure token usage per session. Trim context files. Use /compact at 70%.'],
        ['Developer friction / resistance', 'High', 'Medium', 'Start with CLAUDE.md + hooks. Add skills incrementally. Don\'t mandate full TAI day 1.'],
        ['Hook false positives', 'Medium', 'Medium', 'Bypass-with-reason protocol. Bypass reasons reviewed by architect. Persistent false positives fixed in hook code.'],
        ['Hook crashes blocking work', 'High', 'Low', 'Crashes treated as failures (conservative). Forces immediate fix. Bypass-with-reason available as escape hatch.'],
        ['Stale sprint context', 'Medium', 'Low', 'Architect owns sprint-current.md update. Part of sprint-init skill.'],
        ['Telemetry data volume in git', 'Medium', 'Low', 'JSONL is append-only text. Months of team telemetry < 100 MB. Git handles text efficiently. Archive old sprints if needed.'],
        ['Over-engineering', 'High', 'Medium', 'Ship minimal TAI. Add skills when a pattern repeats 3+ times. Resist framework creep.'],
    ]
)

add_callout(
    'TAI must remain a lightweight discipline layer, not a framework. If the .tai/ directory grows '
    'beyond ~40 files (excluding telemetry data) or the CORE.md exceeds 200 lines, it is too complex. '
    'The test is: can a new team member understand the full TAI setup in under 30 minutes? If not, simplify.',
    'Critical Risk \u2014 Over-Engineering:'
)

# ============================================================
# 12. STRATEGIC VALUE
# ============================================================
add_heading('12. Strategic Value')

doc.add_paragraph(
    'TAI addresses a problem that every AI-assisted development team faces: how to maintain engineering '
    'discipline when AI tools accelerate code generation beyond the team\'s ability to manually verify '
    'consistency. This is not a project-specific problem \u2014 it is an organizational problem.'
)

add_heading('12.1 Broader Applicability', level=2)

applicability = [
    ('Project-portable.', 'The .tai/ directory structure, role system, and skill library can be adapted to any project. Stack-specific skills are swappable; the framework is stack-agnostic.'),
    ('Onboarding accelerator.', 'New team members get full project context, established patterns, and verification guardrails on their first git pull. Time-to-productive drops from weeks to days.'),
    ('Audit trail.', 'Skills, hooks, telemetry, and slash command queries provide a documented record of how AI was used in development \u2014 addressing compliance and governance requirements for AI-assisted engineering.'),
    ('Methodology IP.', 'A mature TAI skill library becomes reusable methodology IP. Best practices for AI-assisted API development, test generation, and architecture review can be packaged and deployed across teams and organizations.'),
]
for label, desc in applicability:
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(desc)

add_heading('12.2 Relationship to AI Adoption', level=2)

doc.add_paragraph(
    'TAI is a concrete implementation of the principles needed for safe, scalable AI-assisted development. '
    'It is the answer to the question every engineering leader will ask: "How do I let my team use AI '
    'without losing control of the codebase?"'
)

doc.add_paragraph(
    'The telemetry and slash command query system make AI-assisted development auditable and transparent '
    '\u2014 a critical requirement for enterprise adoption. When leadership asks "how do you know the AI '
    'isn\'t introducing problems?", the telemetry trail and enforcement hooks are the answer.'
)

# ============================================================
# 13. ACCEPTANCE CRITERIA
# ============================================================
doc.add_page_break()
add_heading('13. Acceptance Criteria')

doc.add_paragraph('TAI v1.0 is considered complete when all of the following are verified:')

criteria = [
    'A team member can clone the project repo, launch Claude Code, and have full project + role + sprint context loaded automatically within 30 seconds.',
    'Role identification works via git email match, with TAI_ROLE env var override functioning correctly.',
    'The architect session and developer session load different context and have access to different skills.',
    'A developer attempting to commit code that violates a configured boundary rule is blocked by the pre-commit hook with a clear error message.',
    'A developer attempting to push code with failing tests is blocked by the pre-push hook.',
    'Hook bypass requires a mandatory reason string. The bypass and reason are logged to telemetry.',
    'If a hook script crashes, the action is blocked (conservative crash policy).',
    'The /new-endpoint slash command produces a complete, pattern-compliant API route with auth, validation, SQL, and tests.',
    'The /pre-pr slash command runs the full verification pipeline and generates a PR description with AI usage disclosure.',
    'Every hook execution writes a structured event to the appropriate JSONL telemetry file.',
    'Telemetry JSONL files are committed to the repo. A GitHub Action aggregates metrics on push.',
    'The following slash commands return correct data: /tai-sessions, /tai-bypasses, /tai-boundaries, /tai-sprint, /tai-health, /tai-usage.',
    'The CLI status line displays: current role, sprint name, ISC progress, hook status, session duration, TAI version.',
    'sprint-current.md accurately reflects the current sprint objectives and is visible in every session.',
    'The .tai/VERSION file contains a valid semver and is displayed in the status line.',
    'The total .tai/ directory is under 40 files (excluding telemetry data) and CORE.md is under 200 lines.',
    'A new team member can understand the full TAI setup in under 30 minutes by reading the README.',
    'All existing project tests still pass with TAI hooks active.',
    'No project-specific references exist in the TAI core framework files (CORE.md, hook scripts, skill templates).',
]
for i, c in enumerate(criteria, 1):
    doc.add_paragraph(f'{i}. {c}')

# ============================================================
# FOOTER
# ============================================================
doc.add_paragraph()
doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('END OF SPECIFICATION')
r.bold = True
r.font.size = Pt(10)

doc.add_paragraph()
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = closing.add_run('TAI Architecture Specification v1.1 \u2014 Jinks Labs \u2014 March 2026')
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

# Save
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'TAI_Architecture_Spec_v1.1.docx')
doc.save(output_path)
print(f'Saved: {output_path}')
